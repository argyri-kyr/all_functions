

import xlsxwriter
import numpy as np
import os
import pandas as pd
import openpyxl
from openpyxl import load_workbook

from globalVars import folder1, folder2, textends
from syllabfindcategory1 import syllabfindcategory1
from syllabfindcategory2 import syllabfindcategory2



def syllabanalysispercategory(sub, text, exam, wordsinitialinfodict, wordsandroi, results_roi, syllaboccurences, flagtext, analysischoice):
    # find categories in syllabes
    category1 = []
    category1_words = []
    category2 = []
    category2_words = []
    category3 = []
    category3_words = []
    category4 = []
    category4_words = []
    
    if( os.path.exists(folder1 + 'texts/' + text + textends + '/syllabAnalysis' + str(analysischoice) + '.xlsx')  and  flagtext==False ):
        textdata = pd.read_excel(folder1 + 'texts/' + text + textends + '/syllabAnalysis' + str(analysischoice) + '.xlsx', header=None)
        textsyllabs = textdata.iloc[:,0].to_list()
    
    if( not os.path.exists(folder1 + 'texts/' + text + textends + '/syllabAnalysis' + str(analysischoice) + '.xlsx') and flagtext==False  ):
        print("You have to make the syllabanalysis in this text first")
        return
    

    if( flagtext==False ):
        workbook = xlsxwriter.Workbook(folder1 + folder2 + 'subject_' + str(sub) + '/all_data/' + text + str(exam) + '/notes/syllabAnalysis' + str(analysischoice) + '.xlsx')
    else:
        workbook =  xlsxwriter.Workbook(folder1 + 'texts/' + text + '_1920/syllabAnalysis' + str(analysischoice) + '.xlsx')
       
       
    worksheet = workbook.add_worksheet()
    
    headers = ["Αριθμος Συλλαβων","#Κοκκινων", "#Συνολου", "Ποσοστο προβλημα/συνολο αυτών", "#Κατηγορια 1", "Ποσοστο κατηγορια1/συνολο", "#Κατηγορια2", "Ποσοστο κατηγορια2/συνολο", "#Κατηγορια3", "Ποσοστο κατηγορια3/συνολο", "Λεξεις ανα κατηγορια" ]
    worksheet.write_row(0, 0, headers)
    worksheet.autofit()

        
    lineno = 1
    cat1_red = []
    cat2_red = []
    cat3_red = []
    for key in results_roi.keys():
        
        if( analysischoice ==1 ):
            category1_temp, category1_words_temp, category2_temp, category2_words_temp, category3_temp, category3_words_temp, category4_temp, category4_words_temp = syllabfindcategory1(results_roi[key], wordsinitialinfodict, wordsandroi)
        elif( analysischoice == 2 ):
            category1_temp, category1_words_temp, category2_temp, category2_words_temp, category3_temp, category3_words_temp, category4_temp, category4_words_temp = syllabfindcategory2(results_roi[key], wordsinitialinfodict, wordsandroi)
        
        category1.append(category1_temp)
        category1_words.append(category1_words_temp)
        category2.append(category2_temp)
        category2_words.append(category2_words_temp)
        category3.append(category3_temp)
        category3_words.append(category3_words_temp)
        category4.append(category4_temp)
        category4_words.append(category4_words_temp)
        
        if( flagtext == False ):
            index = textsyllabs.index(key)
            
            if( textdata.iloc[index, 4] == 0 ):
                pososto1 = -1
            else:
                pososto1 = (len(category1_temp) / textdata.iloc[index, 4])*100
            
            if( textdata.iloc[index, 6] == 0 ):
                pososto2 = -1
            else:
                pososto2 = (len(category2_temp) / textdata.iloc[index, 6])*100
        
            if( textdata.iloc[index, 8] == 0 ):
                pososto3 = -1
            else:
                pososto3 = (len(category3_temp) / textdata.iloc[index, 8])*100
       
        else:
            pososto1 = -1
            pososto2 = -1
            pososto3 = -1

        
        worksheet.write(lineno, 0, key)
        worksheet.write(lineno, 1, len(results_roi[key]))
        worksheet.write(lineno, 2, syllaboccurences[key])
        worksheet.write(lineno, 3, np.round( (len(results_roi[key])/syllaboccurences[key])*100, 2) )
        worksheet.write(lineno, 4, len(category1_temp))
        worksheet.write(lineno, 5, pososto1)
        worksheet.write(lineno, 6, len(category2_temp))
        worksheet.write(lineno, 7, pososto2)
        worksheet.write(lineno, 8, len(category3_temp))
        worksheet.write(lineno, 9, pososto3)
        worksheet.write_row(lineno, 10, category1_words_temp + category2_words_temp + category3_words_temp)
        # worksheet.write_row(lineno, 10, category1_temp + category2_temp + category3_temp)
        # worksheet.write_row(lineno, 10, category3_temp)
        lineno += 1
        
        cat1_red.append( len(category1_temp) )
        cat2_red.append( len(category2_temp) )
        cat3_red.append( len(category3_temp) )
        
        
    find_extra_pososto_per_category(text, flagtext, analysischoice, lineno, workbook, worksheet, cat1_red, cat2_red, cat3_red)
        
    return
    
    

"""
H sinartisi afti briskei ena extra pososto stis gia ta exis: 'Ποσοστο κατηγορια1/συνολο, 'Ποσοστο κατηγορια2/συνολο', 'Ποσοστο κατηγορια3/συνολο', 
xrisimo sthn analysh twn subejct alla gia thn analysi twn keimenwn pithano na min xreazetai, kathws bgazei 100, 100, -1 se ola
opote mallon den mas leei kati stis analysis keimenwn kathws profanws den exw kai reddwelltimewords (paradoxi: exw to arxeio reddwelltimewords.txt alla einai olo to keimeno)
"""
def find_extra_pososto_per_category(text, flagtext, analysischoice, lineno, workbook, worksheet, cat1_red, cat2_red, cat3_red):
    if (flagtext==True):
        # Save current workbook in order to load and write extra pososto
        workbook.close() 
        
    # ---- Read the data ----
    file_path = folder1 + 'texts/' + text + textends + '/syllabAnalysis' + str(analysischoice) + '.xlsx'
    dataall = pd.read_excel(file_path, sheet_name="Sheet1")

    cat1_total = dataall.iloc[:,4].dropna().to_list()
    cat2_total = dataall.iloc[:,6].dropna().to_list()
    cat3_total = dataall.iloc[:,8].dropna().to_list()

    lista1 = [sum(cat1_total), sum(cat2_total), sum(cat3_total)]
    lista2 = [cat1_red, cat2_red, cat3_red]
    lista3 = [5, 7, 9]    
    
    if flagtext==False:
        for j in range(len(lista1)):
            if( lista1[j] == 0 ):
                worksheet.write(lineno+3, lista3[j], -1)
            else:
                pososto = (sum(lista2[j])/lista1[j])*100
                worksheet.write(lineno+3, lista3[j], pososto)
                
        workbook.close() 
    
    
    
    if(flagtext==True):
        # ---- Load workbook and worksheet (existing Excel) ----
        workbook = load_workbook(file_path)
        ws = workbook["Sheet1"]

        # ---- Write into the sheet ----
        # openpyxl uses 1-based indexing -> add +1 to convert
        for j in range(len(lista1)):
            excel_col = lista3[j] + 1      # convert 0-based → 1-based Excel column
            excel_row = lineno + 3 + 1     # convert to Excel row index

            if lista1[j] == 0:
                ws.cell(row=excel_row, column=excel_col, value=-1)
            else:
                pososto = (sum(lista2[j]) / lista1[j]) * 100
                ws.cell(row=excel_row, column=excel_col, value=pososto)

        # ---- Save ----
        workbook.save(file_path)
        
        
    
    
        
    

