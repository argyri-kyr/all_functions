

import os
import pandas as pd
from openpyxl import load_workbook
import xlsxwriter
import operator as op
import numpy


import os,sys,inspect

currentdir = os.path.dirname(os.path.abspath(inspect.getfile(inspect.currentframe())))
parentdir = os.path.dirname(currentdir)
sys.path.insert(0,parentdir)
from auto_adjust_column_width_xlsx import auto_adjust_column_width_xlsx
from export_to_pdf import export_to_pdf


def get_exercise_block(filepath, sheet_name, exercise):
    wb = load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]

    # 1. Find the title cell
    title_row = None
    title_col = None

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and exercise in cell.value:
                title_row = cell.row
                title_col = cell.column
                break
        if title_row:
            break

    if title_row is None:
        raise ValueError(f"Exercise '{exercise}' not found")

    data = []

    for r in range(title_row + 1, ws.max_row + 1):
        first_cell = ws.cell(row=r, column=title_col).value

        # STOP at the NEXT title (any exercise)
        if isinstance(first_cell, str) and "(" in first_cell:
            break

        row_values = [
            ws.cell(row=r, column=c).value
            for c in range(title_col, ws.max_column + 1)
        ]

        # trim trailing empty columns
        while row_values and row_values[-1] is None:
            row_values.pop()

        # keep only meaningful rows
        if any(v is not None for v in row_values):
            data.append(row_values)

    return pd.DataFrame(data)


def get_result(exercise, df):
    last_row = [df.iloc[-1-i] for i in range(len(df)) if df.iloc[-1-i].notna().any()][0]
   
    if(exercise=='Decode Words Pseudo'):
        last_row = df.iloc[-2]
        for cell in last_row:
            if not isinstance(cell, (int, float)):
                last_row = df.iloc[-1]
                break
                
    if(exercise=='Raven Score'):
        last_row = df.iloc[-2]
    
    result = []
    for cell in last_row:
        if cell is None or pd.isna(cell):
            continue

        # Case 1: "3,3"
        if isinstance(cell, str) and "," in cell:
            parts = cell.split(",")
            if all(p.strip().isdigit() for p in parts):
                result.append([int(p) for p in parts])

        # Case 2: extra specific case only for Raven Score and more specific only for True Age cell in order to keep the float part
        elif isinstance(cell, float) and exercise=='Raven Score' and isinstance(cell, float):
            result.append([float(cell)])
            
        # Case 3: single number
        elif isinstance(cell, (int, float, numpy.int64)):
            result.append([int(cell)])
            
        # Case 4: numeric string "3"
        elif isinstance(cell, str) and cell.strip().isdigit():
            result.append([int(cell)])
        
            
    return result


def write_final_comparison_results(subno, score1, score2, output_filepath):
    workbook = xlsxwriter.Workbook(output_filepath)
    worksheet = workbook.add_worksheet()

    worksheet.write(0, 0, "Subject")
    worksheet.write(0, 1, "CHERSO1 Score")
    worksheet.write(0, 2, "CHERSO2 Score")

    worksheet.write(1, 0, subno)
    worksheet.write(1, 1, str(score1))
    worksheet.write(1, 2, str(score2))

    workbook.close()


def create_new_beautify_from_existing(input_filepath, output_filepath, sheet_name):
    # Load source workbook (read-only)
    wb_src = load_workbook(input_filepath, data_only=True)
    ws_src = wb_src.active

    # Create new workbook using xlsxwriter
    wb_new = xlsxwriter.Workbook(output_filepath)
    ws_new = wb_new.add_worksheet(sheet_name)

    header_format = wb_new.add_format({'bold': True})
    
    ws_new.write(0, 0, 'Test', header_format)
    ws_new.write(0, 1, 'Score 1', header_format)
    ws_new.write(0, 2, 'Score 2', header_format)
    
    # Copy data BUT one column right to leave space (one column left) for first radar results
    for r, row in enumerate(ws_src.iter_rows(values_only=True)):
        if( r < 14):
            for c, value in enumerate(row):
                
                if((r==0 and value!=None) and (op.contains(value,'Raven Age') or op.contains(value,'True Age'))):
                    continue
                
                # Fix a bit raven result, instead of showing -1/36, just show -1, means either case that no raven test was done
                if(value=='-1/36'):
                    value='-1'

                if(c==0):
                    ws_new.write(r+1, c, value)
                else:
                    ws_new.write(r+1, c+1, value)
    
    return wb_new, ws_new



def cherso2_score_comparison_cherso_main():
    sheet_name = "ExecutiveSheet"
    exercises_cherso1 = {'Raven Score':0, 'Memory Short Audit Verb':0, 'Memory Work Audit Verb':21, 'Phonological Rhymes':5, 'Phonological Syllables':5, 'Phonological Phonemes':5, 'Decode Words Pseudo':24, 'Visual Orthographic Rec':10, 'Decode Sentences':5}

    sub_start = 13504
    sub_end = 13504
    
    basedir = 'N:/' 
    results_dir = "COMPARISON_RESULTS/"
    for subno in range(sub_start, sub_end + 1):    
        basepath1 = basedir + "CHERSO1/executive_sheets/"
        basepath2 = basedir + "CHERSO2/subject_" + str(subno) + "/"
        
        if not os.path.exists(basepath1) or not os.path.exists(basepath2):
            print(f"Base path for subject {subno} does not exist. Skipping.\n")
            continue
        
        beautify_input_filepath = basepath2 + 'beautify_executive_sheet_' + str(subno) + '.xlsx'
        beautify_output_filepath = basedir + results_dir + "comparison_results_" + str(subno) + ".xlsx"
        new_sheet_name = 'Student Card'
        wb_new, ws_new = create_new_beautify_from_existing(beautify_input_filepath, beautify_output_filepath, new_sheet_name)

        final_scores = []
        raven_extra_info = []
        for exercise in exercises_cherso1.keys():
            
            filepath = basepath1 + 'executive_sheet_' + str(subno) + '.xlsx'
            df = get_exercise_block(filepath, sheet_name, exercise)
            results = get_result(exercise, df)
            
            if(results==[]):
                final_scores.append(str(-1)+ '/' + str(exercises_cherso1[exercise]))
                if(exercise=='Memory Short Audit Verb' or exercise=='Decode Words Pseudo'):
                    final_scores.append(str(-1)+ '/' + str(exercises_cherso1[exercise]))
                continue
                
            if(exercise=='Raven Score'):
                raven_extra_info.append('R_A: ' + str(results[1][0])) # raven age
                raven_extra_info.append('T_A: ' + str(results[2][0])) # true age
                raven_extra_info.append('C_S: ' + str(results[3][0])) # correct score
                final_scores.append(str(results[0][0]))
                
            elif(exercise=='Memory Short Audit Verb'):
                results_new = [results[i][1] for i in range(len(results))]
                results_new_pt1 = results_new[:5]
                results_new_pt2 = results_new[5:]
                final_scores.append(str(sum(results_new_pt1)) +'/' + str(25))
                final_scores.append(str(sum(results_new_pt2)) + '/' + str(13))
                
            elif(exercise=='Memory Work Audit Verb'):
                results_new = [results[i][1] for i in range(len(results))]
                final_scores.append(str(sum(results_new)) + '/' + str(exercises_cherso1[exercise]))
                
            elif(exercise=='Phonological Syllables'):
                given_answer = df.iloc[0].tolist()
                given_answer_pt2 = [int(x) for x in given_answer[5:]]
                final_scores.append(str(sum(given_answer_pt2)) + '/' + str(exercises_cherso1[exercise]))
            
            elif(exercise=='Phonological Phonemes'):
                results_new = results[0][1]
                final_scores.append(str(results_new) + '/' + str(exercises_cherso1[exercise]))
                
            elif(exercise=='Decode Words Pseudo'):
                results_new_pt1 = results[0][0]
                results_new_pt2 = results[1][0]
                final_scores.append(str(results_new_pt1) +'/' + str(exercises_cherso1[exercise])) 
                final_scores.append(str(results_new_pt2) + '/' + str(exercises_cherso1[exercise]))
                
            elif(exercise=='Visual Orthographic Rec'):
                given_answer = df.iloc[1].tolist()
                # about answer1
                given_answer_pt1 = [int(x) for x in given_answer[:5]]
                correct_answer1 = [1,2,4,2,3]
                score_new1 = sum([1 for i in range(5) if given_answer_pt1[i]==correct_answer1[i]])
                final_scores.append(str(score_new1) + '/' + str(5))
                # about answer2
                given_answer_pt2 = [int(x) for x in given_answer[5:]]
                correct_answer2 = [3,2,1,3,4]
                score_new2 = sum([1 for i in range(5) if given_answer_pt2[i]==correct_answer2[i]])
                final_scores.append(str(score_new2) + '/' + str(5))
            
            else:
                results_new = results[0]
                final_scores.append(str(sum(results_new)) + '/' + str(exercises_cherso1[exercise]))


        ws_new.write_column(1, 1, final_scores)
        ws_new.write_row(1, 3, raven_extra_info)
        wb_new.close()
        auto_adjust_column_width_xlsx(beautify_output_filepath, new_sheet_name)
        export_to_pdf(beautify_output_filepath, basedir + results_dir, "comparison_results_" + str(subno) + ".pdf")
        print()
        
        
        
        
if __name__ == "__main__":
    cherso2_score_comparison_cherso_main()